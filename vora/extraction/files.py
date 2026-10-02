"""Downloadable files: classification, discovery, relevance and parsing.

Every file a page links to (or a search returns) is recorded as a reference.
Only formats that can be read as data are *extractable*, and extraction works
on bytes already in memory: nothing is written to disk. Programs are
recorded but never downloaded.

Parsing strategy per format:

* CSV / TSV  -> ``datasets.parse_csv`` (window and place filtering)
* XLSX       -> each sheet becomes a table -> ``datasets.parse_matrix``
* JSON       -> the largest list of records -> ``datasets.parse_matrix``
* PDF        -> text and tables rendered as simple HTML, then the same page
                extractor used for web pages (tables, captions, prose)
"""

from __future__ import annotations

import hashlib
import html as html_lib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from types import MappingProxyType
from urllib.parse import parse_qsl, unquote, urljoin, urlparse

from bs4 import BeautifulSoup

from vora.browser.contracts import ExecutionResult
from vora.shared.contracts import Observation

from vora.extraction.datasets import DatasetLink, parse_csv, parse_matrix
from vora.extraction.semantics import content_tokens, tokens
from vora.extraction.temporal import iso_date

FILE_TYPES: dict[str, frozenset[str]] = {
    "data": frozenset({"csv", "tsv"}),
    "spreadsheet": frozenset({"xlsx", "xlsm", "xls", "ods"}),
    "document": frozenset({"pdf", "docx", "doc", "pptx", "ppt", "odt", "rtf", "epub"}),
    "structured": frozenset({"json", "geojson", "xml", "jsonl", "ndjson"}),
    "archive": frozenset({"zip", "gz", "tgz", "tar", "7z", "rar", "bz2"}),
    "media": frozenset({"png", "jpg", "jpeg", "gif", "svg", "webp", "mp4", "mp3", "wav", "mov", "avi", "webm"}),
    "program": frozenset({"exe", "msi", "apk", "dmg", "bat", "cmd", "sh", "ps1", "jar", "deb", "rpm",
                          "scr", "vbs", "app", "bin", "com", "iso"}),
}
_EXTENSION_TYPE = {extension: kind for kind, extensions in FILE_TYPES.items() for extension in extensions}
# Formats VORA can read as data, and those it reads automatically during a run.
EXTRACTABLE = frozenset({"csv", "tsv", "xlsx", "xlsm", "pdf", "json", "geojson"})
AUTOMATIC = frozenset({"csv", "tsv"})
MAX_FILES_PER_RUN = 50
MAX_PDF_PAGES = 30
MAX_SHEETS = 8
MAX_SHEET_ROWS = 50_000

class FileParseError(ValueError):
    pass


def file_extension(url: str) -> str | None:
    """The file format a URL points to, from its path or a ``format=`` parameter."""
    parsed = urlparse(url)
    query = {key.lower(): value.lower() for key, value in parse_qsl(parsed.query)}
    for key in ("format", "fmt", "type", "output"):
        if query.get(key) in _EXTENSION_TYPE:
            return query[key]
    match = re.search(r"\.([a-z0-9]{2,7})$", unquote(parsed.path).lower())
    return match.group(1) if match and match.group(1) in _EXTENSION_TYPE else None


def file_type(extension: str | None) -> str:
    return _EXTENSION_TYPE.get(extension or "", "other")


def file_id(url: str) -> str:
    return hashlib.sha256(url.encode()).hexdigest()[:16]


@dataclass(frozen=True, slots=True)
class FoundFile:
    url: str
    name: str
    title: str
    extension: str
    file_type: str
    linked_from: str | None
    found_via: str  # page | search | chart

    @property
    def extractable(self) -> bool:
        return self.extension in EXTRACTABLE


def found_file(url: str, title: str, linked_from: str | None, found_via: str,
               extension: str | None = None) -> FoundFile | None:
    extension = extension or file_extension(url)
    if not extension:
        return None
    name = unquote(urlparse(url).path.rsplit("/", 1)[-1]) or url
    return FoundFile(url.split("#", 1)[0], name[:160], " ".join(title.split())[:200], extension,
                     file_type(extension), linked_from, found_via)


def find_file_links(soup: BeautifulSoup, page_url: str, limit: int = 200) -> list[FoundFile]:
    """Every downloadable file a rendered page links to or embeds."""
    found: dict[str, FoundFile] = {}
    for node in soup.select("a[href], iframe[src], embed[src], object[data], source[src]"):
        raw = node.get("href") or node.get("src") or node.get("data") or ""
        url = urljoin(page_url, raw)
        if urlparse(url).scheme not in {"http", "https"}:
            continue
        label = node.get_text(" ", strip=True) or node.get("title") or node.get("aria-label") or ""
        item = found_file(url, label, page_url, "page")
        if item and item.url not in found:
            found[item.url] = item
            if len(found) >= limit:
                break
    return list(found.values())


def from_dataset_link(link: DatasetLink) -> FoundFile:
    """Register a resolved chart download (e.g. an Our World in Data chart CSV)."""
    return FoundFile(link.url, file_name(link.url), link.title, "csv", "data", link.page_url,
                     "chart" if link.kind == "chart" else "page")


def file_name(url: str) -> str:
    return unquote(urlparse(url).path.rsplit("/", 1)[-1]) or url


def relevance(item: FoundFile, terms: list[str], page_title: str = "") -> float:
    """Share of the plan's terms found in the file's title, name or linking page (0–1)."""
    wanted = {token for term in terms for token in content_tokens(term)}
    if not wanted:
        return 0.5
    enough = min(len(wanted), 3)
    own = min(1.0, len(wanted & set(tokens(f"{item.title} {item.name}"))) / enough)
    # The linking page is weak evidence: every file on a relevant page shares it.
    page = min(1.0, len(wanted & set(tokens(page_title))) / enough)
    score = 0.75 * own + 0.15 * page + (0.1 if item.extractable else 0.0)
    return round(min(1.0, score), 3)


# ----------------------------------------------------------------------------
# Parsing (bytes already in memory)
# ----------------------------------------------------------------------------

@dataclass(slots=True)
class FileParse:
    observations: list[Observation] = field(default_factory=list)
    rows_total: int = 0
    note: str = ""


def _sniff(content: bytes, extension: str) -> None:
    head = content[:512].lstrip().lower()
    if head.startswith((b"<!doctype html", b"<html")) and extension not in {"xml"}:
        raise FileParseError("The link returned a web page, not a file")
    if extension == "pdf" and not content.startswith(b"%PDF"):
        raise FileParseError("The file is not a PDF")
    if extension in {"xlsx", "xlsm"} and not content.startswith(b"PK"):
        raise FileParseError("The file is not an Excel workbook")


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _observations(rows: list[dict[str, str]], method: str, *, url: str, title: str, context: str,
                  published_at: str | None, modified_at: str | None, fetched_at: str) -> list[Observation]:
    from vora.extraction.parser import dataset_observations  # the parser imports this package's siblings

    return dataset_observations(rows, url, title=title, context=context, modified_at=modified_at,
                                fetched_at=fetched_at, method=method, published_at=published_at)


def _header_row(matrix: list[list[str]]) -> tuple[int, str]:
    """Index of the header row in a sheet and the title lines above it."""
    titles = []
    for index, row in enumerate(matrix[:30]):
        cells = [cell for cell in row if cell]
        if len(cells) == 1 and index + 1 < len(matrix):
            titles.append(cells[0])
            continue
        if len(cells) >= 2:
            return index, " ".join(titles)[:240]
    return 0, " ".join(titles)[:240]


def _xlsx(content: bytes, **options) -> tuple[list[tuple[str, list[str], list[list[str]]]], dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    properties = workbook.properties
    tables = []
    for sheet in workbook.worksheets[:MAX_SHEETS]:
        matrix = []
        for row in sheet.iter_rows(values_only=True):
            matrix.append(["" if cell is None else str(cell).strip() for cell in row])
            if len(matrix) >= MAX_SHEET_ROWS:
                break
        matrix = [row for row in matrix if any(row)]
        if len(matrix) < 2:
            continue
        header_index, titles = _header_row(matrix)
        context = " · ".join(part for part in (titles, sheet.title) if part)
        tables.append((context, matrix[header_index], matrix[header_index + 1:]))
    workbook.close()
    meta = {"title": getattr(properties, "title", None),
            "created": getattr(properties, "created", None), "modified": getattr(properties, "modified", None)}
    return tables, meta


def _json_records(content: bytes) -> list[dict]:
    payload = json.loads(_decode(content))
    candidates: list[list] = []

    def visit(value, depth: int) -> None:
        if isinstance(value, list) and value and all(isinstance(item, dict) for item in value[:20]):
            candidates.append(value)
        if depth < 3:
            children = value.values() if isinstance(value, dict) else value if isinstance(value, list) else []
            for child in list(children)[:50]:
                visit(child, depth + 1)

    visit(payload, 0)
    if not candidates:
        raise FileParseError("The JSON is not tabular (no list of records found)")
    return max(candidates, key=len)


def _pdf_html(content: bytes) -> tuple[str, dict, int]:
    import pdfplumber

    parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        meta = dict(pdf.metadata or {})
        pages = len(pdf.pages)
        for number, page in enumerate(pdf.pages[:MAX_PDF_PAGES], start=1):
            text = page.extract_text() or ""
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            heading = next((line for line in lines[:3] if 3 <= len(line) <= 120), f"Page {number}")
            parts.append(f"<section><h2>{html_lib.escape(heading)}</h2>")
            for found in page.find_tables():
                rows = [[html_lib.escape(str(cell or "").strip()) for cell in row] for row in found.extract() if row]
                if len(rows) < 2:
                    continue
                # The line just above a table is usually its caption ("Table 1: …").
                top = found.bbox[1]
                above = page.crop((0, max(0, top - 36), page.width, max(1, top))).extract_text() or ""
                caption = next((line.strip() for line in reversed(above.splitlines()) if line.strip()), "")
                head = "".join(f"<th>{cell}</th>" for cell in rows[0])
                body = "".join("<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in rows[1:])
                caption_html = f"<caption>{html_lib.escape(caption)}</caption>" if caption else ""
                parts.append(f"<table>{caption_html}<tr>{head}</tr>{body}</table>")
            # Rejoin hard-wrapped lines into paragraphs for sentence extraction.
            paragraph: list[str] = []
            for line in lines:
                paragraph.append(line)
                if line.endswith((".", ":", ";")) or len(line) < 40:
                    parts.append(f"<p>{html_lib.escape(' '.join(paragraph))}</p>")
                    paragraph = []
            if paragraph:
                parts.append(f"<p>{html_lib.escape(' '.join(paragraph))}</p>")
            parts.append("</section>")
    return "".join(parts), meta, pages


def _pdf_date(value: object) -> str | None:
    text = str(value or "")
    match = re.match(r"D?:?(\d{4})(\d{2})(\d{2})", text)
    return f"{match.group(1)}-{match.group(2)}-{match.group(3)}" if match else iso_date(text)


def parse_file(content: bytes, extension: str, *, url: str, title: str, window: tuple[str, str] | None,
               places: list[str], fetched_at: str | None = None,
               modified_at: str | None = None) -> FileParse:
    """Turn a downloaded file into observations, entirely in memory."""
    if extension not in EXTRACTABLE:
        raise FileParseError(f"Extraction of .{extension} files is not supported yet")
    _sniff(content, extension)
    fetched_at = fetched_at or datetime.now(UTC).isoformat()
    common = {"url": url, "fetched_at": fetched_at}

    if extension in {"csv", "tsv"}:
        table = parse_csv(_decode(content), window=window, places=places)
        rows = _observations(table.rows, "linked_dataset", title=title, context=title,
                             published_at=None, modified_at=modified_at, **common)
        return FileParse(rows, table.total, f"{table.kept:,} of {table.total:,} rows kept")

    if extension in {"json", "geojson"}:
        from vora.extraction.parser import _flatten

        records = [_flatten(record) for record in _json_records(content)[:MAX_SHEET_ROWS]]
        columns = list(dict.fromkeys(key for record in records[:500] for key in record))[:40]
        table = parse_matrix(columns, [[record.get(column, "") for column in columns] for record in records],
                             window=window, places=places)
        rows = _observations(table.rows, "structured_data", title=title, context=title,
                             published_at=None, modified_at=modified_at, **common)
        return FileParse(rows, table.total, f"{table.kept:,} of {table.total:,} records kept")

    if extension in {"xlsx", "xlsm"}:
        tables, meta = _xlsx(content)
        document_title = meta.get("title") or title
        created = meta.get("created").date().isoformat() if meta.get("created") else None
        modified = meta.get("modified").date().isoformat() if meta.get("modified") else modified_at
        result, total, kept = FileParse(), 0, 0
        for context, header, body in tables:
            table = parse_matrix(header, body, window=window, places=places)
            total, kept = total + table.total, kept + table.kept
            result.observations += _observations(
                table.rows, "spreadsheet", title=document_title, context=context or document_title,
                published_at=created, modified_at=modified, **common)
        result.rows_total = total
        result.note = f"{kept:,} of {total:,} rows kept from {len(tables)} sheet(s)"
        return result

    # PDF: render to simple HTML and reuse the web page extractor.
    from vora.extraction.parser import parse_page

    body, meta, pages = _pdf_html(content)
    if not re.search(r"[A-Za-z]{3}", BeautifulSoup(body, "html.parser").get_text(" ")):
        raise FileParseError("No readable text (the PDF may contain scanned images)")
    document_title = str(meta.get("Title") or "").strip() or title
    head = [f"<title>{html_lib.escape(document_title)}</title>"]
    if _pdf_date(meta.get("CreationDate")):
        head.append(f'<meta name="dcterms.created" content="{_pdf_date(meta.get("CreationDate"))}">')
    if _pdf_date(meta.get("ModDate")):
        head.append(f'<meta name="dcterms.modified" content="{_pdf_date(meta.get("ModDate"))}">')
    document = f"<html><head>{''.join(head)}</head><body><main>{body}</main></body></html>"
    rendered = ExecutionResult(
        execution_id=f"file-{file_id(url)}", requested_url=url, final_url=url, title=document_title,
        html=document, status=200, elapsed_seconds=0.0, network_idle_reached=True,
        metadata=MappingProxyType({"fetched_at": fetched_at, "response_headers": MappingProxyType({})}),
    )
    renames = {"html_table": "pdf_table", "text_statement": "pdf_text", "undated_statement": "pdf_text"}
    observations = [
        item.model_copy(update={"method": renames[item.method]})
        for item in parse_page(rendered).observations if item.method in renames
    ]
    read = min(pages, MAX_PDF_PAGES)
    note = f"{len(observations):,} observations from {read} of {pages} page(s)"
    return FileParse(observations, len(observations), note)
